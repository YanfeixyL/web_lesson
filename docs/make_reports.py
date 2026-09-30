from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path

OUT = Path(__file__).parent

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]; r = p.add_run(str(text)); r.bold = bold
    if color: r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(3); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
def base_doc(title):
    d = Document(); sec = d.sections[0]; sec.top_margin = Cm(2.1); sec.bottom_margin = Cm(2.0); sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    normal = d.styles['Normal']; normal.font.name = 'Microsoft YaHei'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); normal.font.size = Pt(10.5); normal.paragraph_format.line_spacing = 1.35
    for name, size, color in [('Title', 22, '1F4E8C'), ('Heading 1', 15, '1F4E8C'), ('Heading 2', 12, '2F6690')]:
        s = d.styles[name]; s.font.name='Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color)
    p=d.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(title)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('《Web 开发技术》平时实验作业'); r.font.color.rgb=RGBColor(120,130,145)
    return d
def h(d, text, level=1): d.add_heading(text, level=level)
def para(d, text, bold_lead=None):
    p=d.add_paragraph();
    if bold_lead and text.startswith(bold_lead): p.add_run(bold_lead).bold=True; p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p
def table(d, headers, rows, widths=None):
    t=d.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,x in enumerate(headers): set_cell_text(t.rows[0].cells[i],x,True,'FFFFFF'); shade(t.rows[0].cells[i],'1F4E8C')
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): set_cell_text(cells[i],x)
    return t

def report():
    d=base_doc('好帮手原型系统设计与实现报告')
    para(d,'本报告说明好帮手互助服务平台的实现范围、运行方式和关键设计。系统围绕“请帮忙”和“我来也”两条业务线展开，普通用户可以发布需求、提交响应并完成确认，管理员可以查看用户和需求状态，平台提供按月份和地域的统计分析。')
    h(d,'一、运行环境配置说明')
    table(d,['项目','配置'],[['操作系统','Windows 10 / Linux'],['运行时','Node.js 20+（本地已用 Node.js 24 验证）'],['前端','Vue 3 CDN、原生 CSS、柱状图组件'],['后台','Express 4、JSON Web Token、Multer'],['数据持久化','开发演示使用 JSON 文件；src/mysql-models.js 提供 Sequelize + MySQL 生产适配'],['数据库','JSON 开发演示；MySQL 8.0 建表脚本见 sql.txt'],['启动命令','cd haobangshou && npm install && npm start'],['访问地址','http://localhost:3000']])
    para(d,'说明：为了让 Windows 10 无需安装数据库也能完成课堂演示，默认使用 data/haobangshou.json 保存数据；接口、实体关系和 sql.txt 已按 MySQL 结构保持一致。部署到 Linux 云服务器时可安装 MySQL，并使用 src/mysql-models.js 的 Sequelize 模型接入。')
    h(d,'二、已实现功能')
    table(d,['模块','实现内容'],[
        ['注册与登录','用户名、姓名、手机号、简介和密码注册；密码长度、数字数量和大小写规则校验；JWT 登录态。'],
        ['个人资料','查看注册信息；仅允许修改手机号、简介和密码，密码仍按注册规则校验。'],
        ['请帮忙','按类型、关键词、地域查询；分页展示；新增、修改、删除自己的未响应需求；支持附件字段。'],
        ['我来也','浏览需求详情；填写响应说明；修改和取消待接受响应；发布者可接受或拒绝，接受后生成成功明细并归档需求。'],
        ['统计分析','按起止月份、地域和类型筛选，列表展示每月发布数、成功响应数与类型明细，柱状图展示趋势并计算转化率。'],
        ['管理员','查看全部用户（密码不返回）和需求状态，账号 admin。'],
        ['交互与异常','新增、删除、接受、拒绝均有确认提示；接口错误以页面 toast 显示；无数据时显示空状态。']])
    h(d,'三、未实现或待完善功能')
    para(d,'1. 当前演示版的图片、视频上传接口已经预留 attachment 字段，页面重点展示文字流程，尚未加入媒体预览组件。')
    para(d,'2. MySQL 生产配置需要在云服务器安装数据库并填写连接参数；默认 JSON 模式用于课堂演示和离线验收。')
    para(d,'3. 管理后台暂未增加复杂的条件组合筛选和导出报表，这些属于作业中的选作扩展。')
    h(d,'四、关键界面与设计说明')
    table(d,['界面','页面说明','对应代码'],[
        ['登录/注册','同一张卡片切换登录和注册；登录页给出两组测试账号，错误信息就地提示。','public/index.html、app.js'],
        ['首页','展示欢迎卡片、三个快捷入口和最近发布的需求卡片。','public/index.html、styles.css'],
        ['需求列表','顶部提供关键词、类型和地域筛选；卡片显示需求主题、描述、地域、发布者与分页。','GET /api/requests'],
        ['需求详情','显示发布者和地域信息；非发布者可提交“我来也”，发布者可处理待接受响应。','GET /api/requests/:id'],
        ['统计分析','筛选条件、统计卡片、月度柱状图和明细表组合展示，避免只显示单一数字。','GET /api/stats'],
        ['个人资料/管理后台','资料页限制可修改字段；管理页分用户和需求两个标签。','PUT /api/users/me、/api/admin/*']])
    h(d,'五、关键实现思路')
    para(d,'系统以用户、地域、需求、响应和成功明细五类数据为核心。需求与响应分别记录发布者和响应者，接受响应时写入 successes，并把需求置为已归档状态。列表接口使用 page、size、offset 参数完成分页；查询条件先在服务端组合，再返回总数和页数。')
    para(d,'前端采用 Vue 3 的响应式状态管理，将登录态保存到 localStorage；每次接口请求自动附加 Bearer Token。组件化拆分需求卡片、空状态和分页器，新增、修改、删除共用同一组弹窗，降低页面重复。')
    table(d,['层次','技术与职责'],[['视图层','Vue 3、HTML、CSS；负责表单、列表、弹窗、图表和确认提示。'],['业务层','Express 路由；负责注册登录、权限校验、需求/响应状态流转和统计聚合。'],['持久化层','JSON 开发适配层；生产部署可按相同实体切换 MySQL，sql.txt 提供完整结构。'],['安全控制','bcryptjs 存储密码摘要；JWT 鉴权；管理员接口二次角色校验；响应不能提交给自己的需求。']])
    h(d,'六、测试记录')
    table(d,['测试项','结果'],[['启动与静态页面','通过，访问 http://localhost:3000 可加载页面。'],['admin 登录','通过，能看到管理后台入口。'],['普通用户登录与注册','通过，手机号和密码规则校验有效。'],['需求分页和筛选','通过，关键词、类型、地域条件均能返回结果。'],['提交响应并接受','通过，响应状态改为已接受，生成 success 记录，需求归档。'],['统计接口','通过，返回月份、发布数、成功数和类型明细。'],['源码检查','通过，node --check src/server.js 与 public/app.js 无语法错误。']])
    h(d,'七、项目文件说明')
    table(d,['文件','用途'],[['src/server.js','Express 服务、认证、业务接口、统计接口和静态文件托管'],['src/models.js','数据结构、JSON 开发持久化与初始化测试数据'],['public/index.html','页面结构和 Vue 模板'],['public/app.js','页面状态、接口调用与交互逻辑'],['public/styles.css','响应式页面样式'],['sql.txt','MySQL 建表及测试数据脚本'],['README.md','运行方式和测试账号']])
    d.save(OUT/'小组实验作业报告.docx')

def progress():
    d=base_doc('好帮手平时实验作业分组及进展情况说明')
    para(d,'本次作业完成一个可运行的好帮手互助服务原型，当前版本已经覆盖课程要求的主要用户流程，并保留了 MySQL 部署入口。')
    h(d,'一、小组成员与分工')
    table(d,['成员','主要工作','贡献比例'],[['严飞（组长）','需求分析、数据模型、后端接口、前端联调、测试与报告整理','100%']])
    h(d,'二、开发环境')
    table(d,['项目','内容'],[['操作系统','Windows 10 开发；Linux 云服务器可部署'],['IDE','Visual Studio Code / 任意 Node.js 编辑器'],['数据库','开发模式 JSON；MySQL 8.0 建表脚本'],['框架','Vue 3、Express、Sequelize（MySQL 适配）、JWT、Multer'],['计划周期','第 14-17 教学周']])
    h(d,'三、作业计划与进展')
    table(d,['阶段','计划','完成情况'],[['需求梳理','提取用户、需求、响应、统计和报告要求','已完成'],['数据模型','设计 users、regions、help_requests、responses、successes 五张表','已完成'],['基础账号','完成注册、登录、密码校验和个人资料','已完成'],['请帮忙模块','分页、筛选、新增、修改、删除和详情','已完成'],['我来也模块','提交、修改、取消、接受、拒绝和成功明细','已完成'],['统计分析','月度、地域、类型筛选，列表与柱状图','已完成'],['测试交付','接口测试、运行说明、sql.txt 和报告','已完成']])
    h(d,'四、当前版本说明')
    para(d,'本地运行 npm install、npm start 后访问 http://localhost:3000。管理员账号为 admin/admin123，普通用户账号为 zhangsan/abc12345。首次启动会自动创建三名测试用户、三个地域、三条需求和两条待响应数据。')
    para(d,'部署时建议使用 Linux 云服务器和 MySQL。先执行 sql.txt，再设置 DB_DIALECT、DB_HOST、DB_PORT、DB_NAME、DB_USER、DB_PASSWORD 和 JWT_SECRET，最后使用 pm2 或 systemd 守护 node src/server.js。')
    h(d,'五、后续补充')
    para(d,'如果继续扩展，可加入图片/视频预览、管理员条件组合查询、月度数据导出和更细的操作日志。目前这些不影响课程要求的核心流程。')
    d.save(OUT/'平时实验作业分组及进展说明.docx')

report(); progress()
